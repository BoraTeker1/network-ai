"""Resume file text extraction (PDF / DOCX).

Local-only, deterministic, no LLM. Turns an uploaded resume file into plain
text that flows into the SAME parsing/upsert path as pasted resume text.

Kept intentionally small: read bytes -> extract text -> hand back a string.
Nothing is persisted to disk; we work entirely in memory.
"""

import io

SUPPORTED_EXTENSIONS = (".pdf", ".docx")
# Generous but sane cap so a bad upload can't exhaust memory (5 MB).
MAX_FILE_BYTES = 5 * 1024 * 1024


class ResumeFileError(Exception):
    """Raised for any user-correctable problem with an uploaded resume file.

    The router maps this to a 400 with the message shown to the user, so keep
    messages clear and actionable.
    """


def _extension(filename: str | None) -> str:
    """Lowercased file extension including the dot, e.g. '.pdf'."""
    if not filename or "." not in filename:
        return ""
    return "." + filename.rsplit(".", 1)[1].lower()


def _extract_pdf(data: bytes) -> str:
    """Extract text from a PDF using pypdf (pure-Python, reliable)."""
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - dependency guard
        raise ResumeFileError(
            "PDF support is not installed on the server (pypdf missing)."
        ) from exc

    try:
        reader = PdfReader(io.BytesIO(data))
        pages = [page.extract_text() or "" for page in reader.pages]
    except Exception as exc:
        raise ResumeFileError(
            "Could not read this PDF. It may be corrupted or password-protected."
        ) from exc
    return "\n".join(pages)


def _extract_docx(data: bytes) -> str:
    """Extract text from a .docx using python-docx (paragraphs + tables)."""
    try:
        import docx
    except ImportError as exc:  # pragma: no cover - dependency guard
        raise ResumeFileError(
            "DOCX support is not installed on the server (python-docx missing)."
        ) from exc

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ResumeFileError(
            "Could not read this DOCX. It may be corrupted or not a real .docx file."
        ) from exc

    parts = [p.text for p in document.paragraphs]
    # Many resumes lay out skills/experience in tables — include those cells too.
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    parts.append(cell.text)
    return "\n".join(parts)


def extract_text_from_file(filename: str | None, data: bytes) -> str:
    """Validate + extract plain text from an uploaded resume file.

    Raises ResumeFileError (-> HTTP 400) for: empty/missing file, unsupported
    extension, oversized file, unreadable file, or no extractable text.
    """
    if not data:
        raise ResumeFileError("No file content received. Please choose a file.")
    if len(data) > MAX_FILE_BYTES:
        mb = MAX_FILE_BYTES // (1024 * 1024)
        raise ResumeFileError(f"File is too large. Maximum size is {mb} MB.")

    ext = _extension(filename)
    if ext not in SUPPORTED_EXTENSIONS:
        raise ResumeFileError(
            "Unsupported file type. Please upload a .pdf or .docx resume."
        )

    text = _extract_pdf(data) if ext == ".pdf" else _extract_docx(data)
    text = text.strip()
    if not text:
        raise ResumeFileError(
            "Couldn't extract any text from this file. If it's a scanned/image "
            "PDF, paste your resume text instead."
        )
    return text
